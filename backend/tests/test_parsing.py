import io
import pytest
from pypdf import PdfWriter
from docx import Document as DocxDocument
from PIL import Image

from app.core.errors import ValidationError
from app.services.parsing.router import get_document_parser_router


@pytest.mark.asyncio
async def test_text_parser():
    router = get_document_parser_router()
    content = b"--- Page 1 ---\nInvoice INV-001\n--- Page 2 ---\nTotal: $1450.00"
    doc = await router.parse_document(content, "test_invoice.txt")
    assert doc.parser == "TextDocumentParser"
    assert len(doc.pages) == 2
    assert "Invoice INV-001" in doc.pages[0].text
    assert "Total: $1450.00" in doc.pages[1].text


@pytest.mark.asyncio
async def test_pdf_parser():
    router = get_document_parser_router()
    # Create valid synthetic PDF using pypdf
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    pdf_bytes_io = io.BytesIO()
    writer.write(pdf_bytes_io)
    pdf_bytes = pdf_bytes_io.getvalue()

    doc = await router.parse_document(pdf_bytes, "invoice.pdf")
    assert doc.parser == "PDFDocumentParser"
    assert len(doc.pages) == 1


@pytest.mark.asyncio
async def test_docx_parser():
    router = get_document_parser_router()
    docx_io = io.BytesIO()
    d = DocxDocument()
    d.add_paragraph("Vendor: Acme Industrial Supplies")
    d.add_paragraph("Total: $1450.00")
    d.save(docx_io)
    docx_bytes = docx_io.getvalue()

    doc = await router.parse_document(docx_bytes, "invoice.docx")
    assert doc.parser == "DocxDocumentParser"
    assert "Vendor: Acme Industrial Supplies" in doc.text
    assert "Total: $1450.00" in doc.text


@pytest.mark.asyncio
async def test_image_parser_valid_png():
    router = get_document_parser_router()
    img_io = io.BytesIO()
    img = Image.new("RGB", (100, 100), color="white")
    img.save(img_io, format="PNG")
    png_bytes = img_io.getvalue()

    doc = await router.parse_document(png_bytes, "scanned_receipt.png")
    assert doc.parser == "ImageOCRParser"
    assert len(doc.pages) == 1


@pytest.mark.asyncio
async def test_corrupted_file_fails_closed():
    router = get_document_parser_router()
    corrupted_pdf = b"THIS IS NOT A VALID PDF CONTENT"
    with pytest.raises(ValidationError):
        await router.parse_document(corrupted_pdf, "broken.pdf")


@pytest.mark.asyncio
async def test_empty_file_fails_closed():
    router = get_document_parser_router()
    with pytest.raises(ValidationError):
        await router.parse_document(b"", "empty.txt")
