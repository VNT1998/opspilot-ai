import io
from typing import List
from docx import Document as DocxDocument
from app.core.errors import ValidationError
from app.services.parsing.result import ParsedDocument, ParsedPage


class DocxDocumentParser:
    """Parser for Microsoft Word (.docx) documents."""

    async def parse(self, content: bytes, filename: str) -> ParsedDocument:
        try:
            doc = DocxDocument(io.BytesIO(content))
        except Exception as e:
            raise ValidationError(f"Invalid or corrupted DOCX file '{filename}': {str(e)}")

        text_lines: List[str] = []

        # Extract text from paragraphs
        for para in doc.paragraphs:
            txt = para.text.strip()
            if txt:
                text_lines.append(txt)

        # Extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    text_lines.append(" | ".join(row_cells))

        full_text = "\n".join(text_lines)
        pages = [ParsedPage(page_number=1, text=full_text, confidence=0.98)]

        return ParsedDocument(
            text=full_text,
            pages=pages,
            parser="DocxDocumentParser",
            parser_version="1.0.0",
            extraction_warnings=[] if full_text.strip() else ["DOCX document contained no text."],
        )
